"""Export web/deck.html to demo/out/deck/plexus-deck.pptx.

Each slide is screenshotted at 1920x1080 (?print=1: no footer, no motion)
and placed full-bleed on a 13.333 x 7.5 in slide, with the slide's
headline as alt text and its <aside class="notes"> as speaker notes. The
HTML stays the source; the .pptx is the leave-behind.

The live slide shows whatever the server last ran. With no results yet,
the policy change is run once (about 20 s of veto calls) for the picture.
Either way v3 is unpublished at the end, so nothing is left behind.

    api/.venv/bin/python demo/export_pptx.py [--url http://127.0.0.1:8766/]
"""

from __future__ import annotations

import argparse
import asyncio
import pathlib

from playwright.async_api import async_playwright
from pptx import Presentation
from pptx.util import Emu, Inches

HERE = pathlib.Path(__file__).resolve().parent
OUT = HERE / "out" / "deck"
SLIDES = OUT / "slides"


async def capture(url: str) -> list[dict]:
    SLIDES.mkdir(parents=True, exist_ok=True)
    async with async_playwright() as p:
        browser = await p.chromium.launch()
        page = await browser.new_page(viewport={"width": 1920, "height": 1080})
        await page.goto(f"{url}deck?print=1#1")
        await page.wait_for_function("window.deckDemo && deckDemo.ready", timeout=30000)
        if not (await page.evaluate("deckDemo.policy().results.jul")):
            print("running the live policy change once for the picture", flush=True)
            await page.evaluate("deckDemo.resetPolicy()")
            await page.evaluate("deckDemo.tryRefund('before').then(() => deckDemo.publish()).then(() => deckDemo.runAfter())")
        await page.evaluate("document.fonts.ready")
        meta = await page.evaluate(
            "[...document.querySelectorAll('.slide')].map(s => ({title: s.dataset.title,"
            " headline: (s.querySelector('h1, h2')?.textContent || '').trim(),"
            " notes: (s.querySelector('.notes')?.textContent || '').trim().replace(/\\s+/g, ' ')}))"
        )
        for i, m in enumerate(meta, 1):
            await page.evaluate(f"deckDemo.goto({i})")
            await asyncio.sleep(1.0)  # live slides fetch on entry
            m["png"] = SLIDES / f"slide-{i:02d}.png"
            await page.screenshot(path=str(m["png"]))
            print(f"  {i:2d}  {m['title']}", flush=True)
        await page.evaluate("deckDemo.resetPolicy()")  # v2 back in effect either way
        await browser.close()
    return meta


def build(meta: list[dict], dest: pathlib.Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    for m in meta:
        slide = prs.slides.add_slide(blank)
        pic = slide.shapes.add_picture(str(m["png"]), Emu(0), Emu(0), prs.slide_width, prs.slide_height)
        pic.name = m["title"]
        pic._element.nvPicPr.cNvPr.set("descr", m["headline"] or m["title"])
        slide.notes_slide.notes_text_frame.text = m["notes"]
    prs.core_properties.title = "Plexus"
    prs.save(str(dest))


async def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--url", default="http://127.0.0.1:8766/")
    args = ap.parse_args()
    meta = await capture(args.url)
    dest = OUT / "plexus-deck.pptx"
    build(meta, dest)
    print(f"wrote {dest} ({len(meta)} slides)")


if __name__ == "__main__":
    asyncio.run(main())
